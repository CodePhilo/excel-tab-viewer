"""
core/exporter.py

Saves a DataFrame (a tab's current view: filtered, sorted, visible columns
only) as a new .xlsx or .csv file. The source workbook is never touched.
"""

from __future__ import annotations

import os

import pandas as pd

MAX_EXPORT_COLUMN_WIDTH = 50  # characters, for the .xlsx column auto-fit


class ExportError(Exception):
    """Raised for any problem writing the export, with a user-friendly message."""


def _clean_sheet_name(name: str) -> str:
    for ch in '[]:*?/\\':
        name = name.replace(ch, "_")
    return name[:31] or "Sheet1"


def _strip_illegal_characters(df: pd.DataFrame) -> pd.DataFrame:
    """openpyxl refuses text containing control characters (they can come
    from .xls files or pasted data); drop just those characters."""
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

    df = df.copy()
    for col in df.columns:
        if df[col].dtype == object:
            df[col] = df[col].map(lambda v: ILLEGAL_CHARACTERS_RE.sub("", v) if isinstance(v, str) else v)
    return df


def export_frame(df: pd.DataFrame, path: str, sheet_name: str = "Sheet1") -> None:
    """Write `df` to `path`; the format follows the extension (.csv, else .xlsx).
    Raises ExportError on failure."""
    try:
        if path.lower().endswith(".csv"):
            # utf-8-sig: the byte-order mark makes Excel open non-English
            # text (Arabic, accents, ...) correctly instead of as garbage.
            df.to_csv(path, index=False, encoding="utf-8-sig")
            return
        df = _strip_illegal_characters(df)
        sheet_name = _clean_sheet_name(sheet_name)
        with pd.ExcelWriter(path, engine="openpyxl") as writer:
            df.to_excel(writer, sheet_name=sheet_name, index=False)
            ws = writer.sheets[sheet_name]
            ws.freeze_panes = "A2"  # keep the header row visible while scrolling
            for idx, col in enumerate(df.columns, start=1):
                sample = [str(col)] + [str(v) for v in df[col].head(200) if pd.notna(v)]
                width = min(max(len(s) for s in sample) + 2, MAX_EXPORT_COLUMN_WIDTH)
                ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = width
            _set_date_formats(ws, df)
            _keep_equals_text_as_text(ws, df)
    except PermissionError:
        raise ExportError(
            f"Could not write '{os.path.basename(path)}'.\n"
            "If it's open in Excel, close it and try again."
        )
    except OSError as exc:
        raise ExportError(f"Could not write '{os.path.basename(path)}'.\nDetails: {exc}")


def _keep_equals_text_as_text(ws, df: pd.DataFrame) -> None:
    """openpyxl writes any text starting with '=' as a formula, so a cell
    that merely reads '=== total' would become a broken formula."""
    for idx, col in enumerate(df.columns, start=1):
        if df[col].dtype != object:
            continue
        for pos, value in enumerate(df[col]):
            if isinstance(value, str) and value.startswith("="):
                ws.cell(row=pos + 2, column=idx).data_type = "s"


def _set_date_formats(ws, df: pd.DataFrame) -> None:
    """pandas writes every datetime as 'YYYY-MM-DD HH:MM:SS'; show date-only
    columns as plain dates, the way the app shows them."""
    for idx, col in enumerate(df.columns, start=1):
        series = df[col]
        if not pd.api.types.is_datetime64_any_dtype(series):
            continue
        values = series.dropna()
        if len(values) and (values == values.dt.normalize()).all():
            for (cell,) in ws.iter_rows(min_row=2, min_col=idx, max_col=idx):
                cell.number_format = "yyyy-mm-dd"
