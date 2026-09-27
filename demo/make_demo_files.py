"""
demo/make_demo_files.py

Builds the sample workbooks in this folder, for trying out and testing the
app without real data. The data is random but seeded, so every run
produces the same values.

    python demo/make_demo_files.py

What each file shows off:
  Sales_2026-01/02/03.xlsx  monthly reports for File > Open Files by Name
                            Pattern (e.g. begins with "Sales_"); numbers,
                            dates, text, blanks, phone numbers with leading
                            zeros, and "N/A" text to sort, filter and total
  Costs_2026.xlsx           a file a "Sales_" pattern should NOT match
  Customers.xlsx            two sheets; multi-line notes, long text
  Inventory_Report.xlsx     title rows above the real header (row 4) and a
                            stray value far off in column XFD, which used
                            to freeze the app
"""

from __future__ import annotations

import datetime
import os
import random

from openpyxl import Workbook
from openpyxl.styles import Font

HERE = os.path.dirname(os.path.abspath(__file__))

REGIONS = ["North", "South", "East", "West", "Central"]
REPS = ["Alice Martin", "Omar Haddad", "Mei Chen", "Carlos Ruiz", "Sara Nilsson", "Yusuf Kaya"]
PRODUCTS = {
    "Laptop 14\"": 899.0,
    "Laptop 16\"": 1299.0,
    "Monitor 27\"": 249.99,
    "Keyboard": 39.5,
    "Mouse": 19.99,
    "Docking Station": 159.0,
    "Headset": 74.25,
    "Webcam": 59.9,
}
STATUSES = ["Delivered", "Delivered", "Delivered", "Shipped", "Pending", "Cancelled"]
CITIES = ["Amsterdam", "Berlin", "Cairo", "Dubai", "Istanbul", "London", "Madrid", "Riyadh", "Toronto"]


def _bold_header(ws, row: int = 1) -> None:
    for cell in ws[row]:
        cell.font = Font(bold=True)


def _phone(rng: random.Random) -> str:
    # Stored as text on purpose: the leading zeros must survive loading.
    return "00" + str(rng.choice([31, 49, 20, 971, 90, 44])) + "".join(str(rng.randint(0, 9)) for _ in range(8))


def make_sales(month: int, rng: random.Random) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Orders"
    ws.append([
        "Order ID", "Order Date", "Region", "Sales Rep", "Product",
        "Quantity", "Unit Price", "Total", "Status", "Customer Phone", "Discount Code",
    ])
    _bold_header(ws)
    for i in range(1, 181):
        product = rng.choice(list(PRODUCTS))
        qty = rng.randint(1, 25)
        price = PRODUCTS[product]
        day = rng.randint(1, 28)
        ws.append([
            f"SO-2026{month:02d}-{i:04d}",
            datetime.date(2026, month, day),
            rng.choice(REGIONS),
            rng.choice(REPS),
            product,
            qty,
            price,
            round(qty * price, 2),
            rng.choice(STATUSES),
            _phone(rng),
            rng.choice(["", "", "", "SPRING10", "VIP", "N/A"]) or None,
        ])
    for column, width in zip("ABCDEFGHIJK", (18, 12, 10, 16, 18, 10, 11, 12, 11, 17, 14)):
        ws.column_dimensions[column].width = width
    wb.save(os.path.join(HERE, f"Sales_2026-{month:02d}.xlsx"))


def make_costs(rng: random.Random) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Costs"
    ws.append(["Month", "Category", "Amount", "Approved By"])
    _bold_header(ws)
    for month in range(1, 13):
        for category in ("Rent", "Salaries", "Travel", "Software"):
            ws.append([
                datetime.date(2026, month, 1),
                category,
                round(rng.uniform(800, 25000), 2),
                rng.choice(REPS),
            ])
    wb.save(os.path.join(HERE, "Costs_2026.xlsx"))


def make_customers(rng: random.Random) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Customers"
    ws.append(["Customer ID", "Name", "City", "Phone", "Since", "Active", "Credit Limit"])
    _bold_header(ws)
    first = ["Lina", "Noah", "Aya", "Lucas", "Zeynep", "Adam", "Hana", "Leo", "Maya", "Karim"]
    last = ["Schmidt", "Rahman", "Jansen", "Yilmaz", "Costa", "Okafor", "Novak", "Dubois"]
    for i in range(1, 61):
        ws.append([
            f"C{i:04d}",
            f"{rng.choice(first)} {rng.choice(last)}",
            rng.choice(CITIES),
            _phone(rng),
            datetime.date(rng.randint(2015, 2025), rng.randint(1, 12), rng.randint(1, 28)),
            rng.random() > 0.2,
            rng.choice([1000, 2500, 5000, 10000, None]),
        ])

    notes = wb.create_sheet("Notes")
    notes.append(["Customer ID", "Date", "Note"])
    _bold_header(notes)
    samples = [
        "Called about a late delivery.\nPromised a follow-up by Friday.",
        "Asked for a quote on 20 docking stations.",
        "Prefers email contact. Invoices go to the finance team, not the buyer.",
        "Returned one monitor (dead pixels). Replacement shipped.\nCase closed.",
        "N/A",
    ]
    for i in range(1, 41):
        notes.append([
            f"C{rng.randint(1, 60):04d}",
            datetime.datetime(2026, rng.randint(1, 3), rng.randint(1, 28), rng.randint(8, 17), rng.choice([0, 15, 30, 45])),
            rng.choice(samples),
        ])
    wb.save(os.path.join(HERE, "Customers.xlsx"))


def make_inventory(rng: random.Random) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Stock"
    ws.append(["Warehouse Inventory Report"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"Generated {datetime.date(2026, 3, 31):%d %B %Y}"])
    ws.append([])
    ws.append(["SKU", "Item", "Warehouse", "On Hand", "Reorder Level", "Unit Cost"])
    _bold_header(ws, 4)
    for i in range(1, 501):
        product = rng.choice(list(PRODUCTS))
        ws.append([
            f"SKU-{i:05d}",
            product,
            rng.choice(["AMS-1", "BER-2", "DXB-1"]),
            rng.randint(0, 400),
            rng.choice([20, 50, 100]),
            round(PRODUCTS[product] * 0.6, 2),
        ])
    # A leftover value at the very last column of the sheet (XFD): opening
    # this file used to build 16,384 columns and freeze the app.
    ws.cell(row=12, column=16384, value="old note")
    wb.save(os.path.join(HERE, "Inventory_Report.xlsx"))


def main() -> None:
    rng = random.Random(2026)
    for month in (1, 2, 3):
        make_sales(month, rng)
    make_costs(rng)
    make_customers(rng)
    make_inventory(rng)
    print(f"Demo files written to {HERE}")


if __name__ == "__main__":
    main()
